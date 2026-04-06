from pypdf import PdfReader
from docx import Document
import io

def extract_pdf_text(file):
    if isinstance(file, bytes):
        file = io.BytesIO(file)

    reader = PdfReader(file)
    text = ""

    for page in reader.pages:
        page_text = page.extract_text()
        if page_text:
            text += page_text + " "

    return text.strip()


def extract_docx_text(file):
    # DOCX ke liye bhi same fix
    if isinstance(file, bytes):
        file = io.BytesIO(file)

    doc = Document(file)
    text = ""

    for para in doc.paragraphs:
        text += para.text + " "

    return text.strip()

def extract_pdf_text(file):
    reader = PdfReader(file)
    return " ".join([p.extract_text() for p in reader.pages])

def extract_docx_text(file):
    doc = Document(file)
    return " ".join([p.text for p in doc.paragraphs])
